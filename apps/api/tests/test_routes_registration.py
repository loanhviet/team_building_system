"""registration_service.submit_registration/assert_can_edit enforced at the
HTTP boundary: terms gate on submit, and edits blocked once the event leaves
registration_open (or its own record was cancelled)."""

import io
from datetime import timedelta

from openpyxl import load_workbook
from sqlalchemy import select

from app.core.time import utcnow
from app.models.bus import Bus, BusAssignment
from app.models.enums import EventStatus
from app.models.flight import Flight, FlightAssignment
from app.models.gala import GalaSeat, GalaTable
from app.models.hotel import Hotel, Room, RoomAssignment
from app.models.notification import EmailOutbox
from app.models.registration import Registration
from app.services.event_service import upsert_setting
from app.services.registration_service import remindable_employees, reminder_dedupe_key
from tests.conftest import make_employee


async def test_submit_without_agreeing_terms_is_rejected(client, world, auth_headers):
    resp = await client.post(
        f"/api/events/{world.event.id}/registrations/me/submit",
        headers=auth_headers(world.employee_user),
        json={"is_participating": True, "agreed_terms": False},
    )
    assert resp.status_code == 400
    assert resp.json()["error"]["code"] == "terms_not_agreed"


async def test_submit_succeeds_when_terms_agreed(client, world, auth_headers):
    resp = await client.post(
        f"/api/events/{world.event.id}/registrations/me/submit",
        headers=auth_headers(world.employee_user),
        json={"is_participating": True, "agreed_terms": True, "terms_version": "v1"},
    )
    assert resp.status_code == 200
    assert resp.json()["status"] == "submitted"


async def test_edit_blocked_after_registration_closed(client, world, auth_headers, db_session):
    world.event.status = EventStatus.registration_closed
    await db_session.commit()

    resp = await client.put(
        f"/api/events/{world.event.id}/registrations/me",
        headers=auth_headers(world.employee_user),
        json={"is_participating": True},
    )
    assert resp.status_code == 400
    assert resp.json()["error"]["code"] == "registration_closed"


async def test_viewing_registration_never_creates_a_draft(client, world, auth_headers, db_session):
    await db_session.delete(world.registration)
    await db_session.commit()

    response = await client.get(
        f"/api/events/{world.event.id}/registrations/me", headers=auth_headers(world.employee_user)
    )
    assert response.status_code == 200
    assert response.json() is None
    rows = (
        await db_session.execute(
            select(Registration).where(Registration.event_id == world.event.id)
        )
    ).scalars().all()
    assert rows == []


async def test_explicit_draft_endpoint_creates_registration(client, world, auth_headers, db_session):
    await db_session.delete(world.registration)
    await db_session.commit()

    response = await client.post(
        f"/api/events/{world.event.id}/registrations/me/draft",
        headers=auth_headers(world.employee_user),
    )
    assert response.status_code == 201
    assert response.json()["status"] == "draft"


async def test_registration_note_and_cancel_reason_have_length_limits(client, world, auth_headers):
    note = await client.put(
        f"/api/events/{world.event.id}/registrations/me",
        headers=auth_headers(world.employee_user),
        json={"wish_note": "x" * 2001},
    )
    assert note.status_code == 422
    reason = await client.post(
        f"/api/events/{world.event.id}/registrations/me/cancel",
        headers=auth_headers(world.employee_user),
        json={"reason": "x" * 501},
    )
    assert reason.status_code == 422


async def test_edit_blocked_before_registration_open_time(
    client, world, auth_headers, db_session
):
    world.event.registration_open_at = utcnow() + timedelta(hours=1)
    await db_session.commit()
    resp = await client.put(
        f"/api/events/{world.event.id}/registrations/me",
        headers=auth_headers(world.employee_user),
        json={"wish_note": "too early"},
    )
    assert resp.status_code == 400
    assert resp.json()["error"]["code"] == "registration_not_open"


async def test_cancel_removes_locked_operational_assignments(
    client, world, auth_headers, db_session
):
    from app.models.event import TransportLeg

    flight = Flight(event_id=world.event.id, flight_code="VN1", direction="outbound", capacity=10)
    leg = TransportLeg(
        event_id=world.event.id, code="L1", name="Ra sân bay", direction="outbound"
    )
    hotel = Hotel(event_id=world.event.id, code="H1", name="Hotel")
    db_session.add_all([flight, leg, hotel])
    await db_session.flush()
    room = Room(hotel_id=hotel.id, room_number="101", capacity=2)
    bus = Bus(event_id=world.event.id, leg_id=leg.id, code="B1", capacity=10)
    db_session.add_all([room, bus])
    await db_session.flush()
    db_session.add_all([
        FlightAssignment(event_id=world.event.id, flight_id=flight.id, employee_id=world.employee.id, direction="outbound", is_locked=True),
        RoomAssignment(event_id=world.event.id, room_id=room.id, employee_id=world.employee.id),
        BusAssignment(event_id=world.event.id, leg_id=leg.id, bus_id=bus.id, employee_id=world.employee.id, is_locked=True),
    ])
    await db_session.commit()

    resp = await client.post(
        f"/api/events/{world.event.id}/registrations/me/cancel",
        headers=auth_headers(world.employee_user), json={"reason": "Không tham gia được"},
    )
    assert resp.status_code == 200
    assert resp.json()["status"] == "cancelled"
    from sqlalchemy import select

    assert (await db_session.execute(select(FlightAssignment))).scalars().all() == []
    assert (await db_session.execute(select(RoomAssignment))).scalars().all() == []
    assert (await db_session.execute(select(BusAssignment))).scalars().all() == []


async def test_opt_out_clears_assignments_and_personal_gala_seat(
    client, world, auth_headers, db_session,
):
    flight = Flight(event_id=world.event.id, flight_code="VN1", direction="outbound", capacity=10)
    table = GalaTable(event_id=world.event.id, code="B1", seat_count=1)
    db_session.add_all([flight, table])
    await db_session.flush()
    seat = GalaSeat(
        table_id=table.id, seat_number=1, status="confirmed",
        team_id=world.team.id, employee_id=world.employee.id,
    )
    db_session.add_all([seat, FlightAssignment(
        event_id=world.event.id, flight_id=flight.id, employee_id=world.employee.id,
        direction="outbound", is_locked=True,
    )])
    await db_session.commit()

    resp = await client.put(
        f"/api/events/{world.event.id}/registrations/me",
        headers=auth_headers(world.employee_user), json={"is_participating": False},
    )
    assert resp.status_code == 200
    assert resp.json()["is_participating"] is False
    assert (await db_session.execute(select(FlightAssignment))).scalars().all() == []
    await db_session.refresh(seat)
    assert seat.employee_id is None
    assert seat.status == "confirmed"


async def test_latest_registration_survives_registration_closed(client, world, auth_headers, db_session):
    # world already has a submitted registration on the registration_open event
    resp = await client.get("/api/registrations/me", headers=auth_headers(world.employee_user))
    assert resp.status_code == 200
    assert resp.json()["id"] == world.registration.id

    world.event.status = EventStatus.registration_closed
    await db_session.commit()

    # GET /events/current would now return null — the unscoped endpoint must
    # still find this registration so the CBNV can see what they submitted
    resp = await client.get("/api/registrations/me", headers=auth_headers(world.employee_user))
    assert resp.status_code == 200
    assert resp.json()["id"] == world.registration.id


async def test_latest_registration_null_when_event_completed(client, world, auth_headers, db_session):
    world.event.status = EventStatus.event_completed
    await db_session.commit()

    resp = await client.get("/api/registrations/me", headers=auth_headers(world.employee_user))
    assert resp.status_code == 200
    assert resp.json() is None


async def test_export_respects_active_filters(client, world, auth_headers, db_session):
    """Regression test for §3.6: export must apply the same filters as the
    on-screen list, not silently dump the whole event."""
    other = await make_employee(db_session, team=world.team, site=world.site, code="NV002")
    db_session.add(
        Registration(
            event_id=world.event.id, employee_id=other.employee.id, status="draft",
            is_participating=None,
        )
    )
    await db_session.commit()

    resp_all = await client.get(
        f"/api/events/{world.event.id}/registrations/export",
        headers=auth_headers(world.organizer_user),
    )
    assert resp_all.status_code == 200
    wb_all = load_workbook(io.BytesIO(resp_all.content))
    assert wb_all.active.max_row - 1 == 2  # header + both registrations

    resp_filtered = await client.get(
        f"/api/events/{world.event.id}/registrations/export?status_filter=submitted",
        headers=auth_headers(world.organizer_user),
    )
    assert resp_filtered.status_code == 200
    wb_filtered = load_workbook(io.BytesIO(resp_filtered.content))
    assert wb_filtered.active.max_row - 1 == 1
    assert wb_filtered.active.cell(row=2, column=1).value == world.employee.employee_code


async def test_admin_list_includes_position_and_searches_code(client, world, auth_headers, db_session):
    world.employee.position = "Backend Eng"
    await db_session.commit()

    resp = await client.get(
        f"/api/events/{world.event.id}/registrations",
        headers=auth_headers(world.organizer_user),
    )
    assert resp.status_code == 200
    row = next(r for r in resp.json() if r["employee_id"] == world.employee.id)
    assert row["position"] == "Backend Eng"
    assert row["team_code"] == world.team.code

    code = world.employee.employee_code
    resp = await client.get(
        f"/api/events/{world.event.id}/registrations?search={code}",
        headers=auth_headers(world.organizer_user),
    )
    assert resp.status_code == 200
    assert any(r["employee_id"] == world.employee.id for r in resp.json())


async def test_remind_unsubmitted_is_admin_only_and_queues(client, world, auth_headers, db_session):
    db_session.add(
        Registration(
            event_id=world.event.id,
            employee_id=(await make_employee(db_session, team=world.team, site=world.site, code="NV009")).employee.id,
            status="draft",
        )
    )
    await db_session.commit()

    denied = await client.post(
        f"/api/events/{world.event.id}/registrations/remind",
        headers=auth_headers(world.employee_user),
    )
    assert denied.status_code == 403

    ok = await client.post(
        f"/api/events/{world.event.id}/registrations/remind",
        headers=auth_headers(world.organizer_user),
    )
    assert ok.status_code == 202
    assert ok.json()["queued"] >= 1


async def test_registration_confirmed_dedupes_by_content_not_just_registration_id(
    client, world, auth_headers, db_session,
):
    """R7 E5: resubmitting with the same content must not re-send; resubmitting
    with *different* content (a different Ca here) must send an updated email —
    the old dedupe key was `registration_confirmed:{reg.id}` alone, so a
    changed resubmission silently never got a fresh confirmation."""
    from sqlalchemy import select

    from app.models.event import Shift
    from app.models.notification import EmailOutbox

    shift2 = Shift(event_id=world.event.id, code="CA2", name="Ca 2", sort_order=2)
    db_session.add(shift2)
    await db_session.commit()

    await client.put(
        f"/api/events/{world.event.id}/registrations/me",
        headers=auth_headers(world.employee_user),
        json={"shift_id": world.shift.id},
    )
    resp = await client.post(
        f"/api/events/{world.event.id}/registrations/me/submit",
        headers=auth_headers(world.employee_user),
        json={"is_participating": True, "agreed_terms": True, "terms_version": "v1"},
    )
    assert resp.status_code == 200

    # identical resubmit — no new outbox row
    resp = await client.post(
        f"/api/events/{world.event.id}/registrations/me/submit",
        headers=auth_headers(world.employee_user),
        json={"is_participating": True, "agreed_terms": True, "terms_version": "v1"},
    )
    assert resp.status_code == 200
    rows = (
        await db_session.execute(
            select(EmailOutbox).where(EmailOutbox.template_code == "registration_confirmed")
        )
    ).scalars().all()
    assert len(rows) == 1

    # resubmit with a different Ca — content changed, a second email is queued
    await client.put(
        f"/api/events/{world.event.id}/registrations/me",
        headers=auth_headers(world.employee_user),
        json={"shift_id": shift2.id},
    )
    resp = await client.post(
        f"/api/events/{world.event.id}/registrations/me/submit",
        headers=auth_headers(world.employee_user),
        json={"is_participating": True, "agreed_terms": True, "terms_version": "v1"},
    )
    assert resp.status_code == 200
    rows = (
        await db_session.execute(
            select(EmailOutbox).where(EmailOutbox.template_code == "registration_confirmed")
        )
    ).scalars().all()
    assert len(rows) == 2


async def test_remind_reaches_employees_who_never_opened_the_form(
    client, world, auth_headers, db_session
):
    """The group that most needs the nudge has no registration row at all — a
    draft only exists once they've visited the page."""
    never_opened = await make_employee(
        db_session, team=world.team, site=world.site, code="NV020"
    )
    drafted = await make_employee(db_session, team=world.team, site=world.site, code="NV021")
    db_session.add(
        Registration(event_id=world.event.id, employee_id=drafted.employee.id, status="draft")
    )
    opted_out = await make_employee(db_session, team=world.team, site=world.site, code="NV022")
    db_session.add(
        Registration(
            event_id=world.event.id, employee_id=opted_out.employee.id, status="cancelled"
        )
    )
    await db_session.commit()

    targets = await remindable_employees(db_session, world.event.id)
    ids = {e.id for e in targets}

    assert never_opened.employee.id in ids
    assert drafted.employee.id in ids
    assert opted_out.employee.id not in ids, "huỷ đăng ký là câu trả lời, không nhắc"
    # world.employee already submitted
    assert world.employee.id not in ids

    res = await client.post(
        f"/api/events/{world.event.id}/registrations/remind",
        headers=auth_headers(world.organizer_user),
    )
    assert res.status_code == 202
    assert res.json()["queued"] == len(targets)

    dashboard = await client.get(
        f"/api/events/{world.event.id}/dashboard", headers=auth_headers(world.organizer_user)
    )
    # the number BTC confirms in the dialog is the number that gets mailed
    assert dashboard.json()["remindable_count"] == len(targets)


async def test_remind_blocked_once_registration_window_is_closed(
    client, world, auth_headers, db_session
):
    await make_employee(db_session, team=world.team, site=world.site, code="NV023")
    world.event.status = EventStatus.event_completed
    await db_session.commit()

    res = await client.post(
        f"/api/events/{world.event.id}/registrations/remind",
        headers=auth_headers(world.organizer_user),
    )
    assert res.status_code == 400
    assert res.json()["error"]["code"] == "registration_not_open"

    dashboard = await client.get(
        f"/api/events/{world.event.id}/dashboard", headers=auth_headers(world.organizer_user)
    )
    assert dashboard.json()["remindable_count"] == 0


async def test_submitted_terms_version_comes_from_the_event_not_the_client(
    client, world, auth_headers, db_session
):
    await upsert_setting(db_session, world.event.id, "terms_version", "v7")
    await db_session.commit()

    res = await client.post(
        f"/api/events/{world.event.id}/registrations/me/submit",
        headers=auth_headers(world.employee_user),
        json={"is_participating": True, "agreed_terms": True, "terms_version": "hacked"},
    )
    assert res.status_code == 200
    assert res.json()["terms_version"] == "v7"


async def test_already_nudged_today_drops_out_of_the_remind_count(
    client, world, auth_headers, db_session
):
    """The count BTC confirms must be what actually gets mailed: the worker
    dedupes per employee per UTC day, so a second click on the same day used to
    still answer "queued: 119" while sending nothing."""
    await make_employee(db_session, team=world.team, site=world.site, code="NV040")
    await db_session.commit()

    targets = await remindable_employees(db_session, world.event.id)
    assert len(targets) >= 1

    # simulate the worker having mailed all but one of them today
    day = utcnow().date().isoformat()
    for employee in targets[:-1]:
        db_session.add(
            EmailOutbox(
                event_id=world.event.id, to_email=employee.email,
                template_code="registration_reminder", payload_json={},
                dedupe_key=reminder_dedupe_key(world.event.id, employee.id, day),
            )
        )
    await db_session.commit()

    res = await client.post(
        f"/api/events/{world.event.id}/registrations/remind",
        headers=auth_headers(world.organizer_user),
    )
    assert res.json()["queued"] == 1

    dashboard = await client.get(
        f"/api/events/{world.event.id}/dashboard", headers=auth_headers(world.organizer_user)
    )
    assert dashboard.json()["remindable_count"] == 1

    # and once the last one is nudged too, the button has nothing left to do
    db_session.add(
        EmailOutbox(
            event_id=world.event.id, to_email=targets[-1].email,
            template_code="registration_reminder", payload_json={},
            dedupe_key=reminder_dedupe_key(world.event.id, targets[-1].id, day),
        )
    )
    await db_session.commit()

    exhausted = await client.post(
        f"/api/events/{world.event.id}/registrations/remind",
        headers=auth_headers(world.organizer_user),
    )
    assert exhausted.status_code == 400
    assert exhausted.json()["error"]["code"] == "nothing_to_remind"
