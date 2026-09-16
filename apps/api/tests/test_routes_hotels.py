"""Hotel/room assignment integrity and correction workflows."""

import io

from openpyxl import Workbook
from sqlalchemy import select

from app.core.time import utcnow
from app.models.enums import EventStatus
from app.models.hotel import Hotel, Room, RoomAssignment, RoomType
from app.models.registration import Registration
from tests.conftest import make_employee


async def _make_hotel_with_room(db_session, event_id: int) -> tuple[Hotel, RoomType, Room]:
    hotel = Hotel(event_id=event_id, name="Khach san A")
    db_session.add(hotel)
    await db_session.flush()
    room_type = RoomType(hotel_id=hotel.id, name="Doi", capacity=2, quantity=10)
    db_session.add(room_type)
    await db_session.flush()
    room = Room(hotel_id=hotel.id, room_type_id=room_type.id, room_number="101", capacity=2)
    db_session.add(room)
    await db_session.flush()
    await db_session.commit()
    return hotel, room_type, room


async def test_delete_room_with_assignment_is_blocked_then_succeeds_after_unassign(
    client, world, auth_headers, db_session
):
    hotel, _room_type, room = await _make_hotel_with_room(db_session, world.event.id)

    resp = await client.post(
        f"/api/events/{world.event.id}/room-assignments/assign",
        headers=auth_headers(world.organizer_user),
        json={"employee_id": world.employee.id, "room_id": room.id},
    )
    assert resp.status_code == 200
    assignment_id = resp.json()["id"]

    resp = await client.delete(
        f"/api/events/{world.event.id}/hotels/{hotel.id}/rooms/{room.id}",
        headers=auth_headers(world.organizer_user),
    )
    assert resp.status_code == 409
    assert resp.json()["error"]["code"] == "room_in_use"

    resp = await client.delete(
        f"/api/events/{world.event.id}/room-assignments/{assignment_id}",
        headers=auth_headers(world.organizer_user),
    )
    assert resp.status_code == 204

    resp = await client.delete(
        f"/api/events/{world.event.id}/hotels/{hotel.id}/rooms/{room.id}",
        headers=auth_headers(world.organizer_user),
    )
    assert resp.status_code == 204


async def test_delete_room_type_in_use_is_blocked(client, world, auth_headers, db_session):
    hotel, room_type, _room = await _make_hotel_with_room(db_session, world.event.id)

    resp = await client.delete(
        f"/api/events/{world.event.id}/hotels/{hotel.id}/room-types/{room_type.id}",
        headers=auth_headers(world.organizer_user),
    )
    assert resp.status_code == 409
    assert resp.json()["error"]["code"] == "room_type_in_use"


async def test_assign_room_rejects_non_participant(client, world, auth_headers, db_session):
    _hotel, _room_type, room = await _make_hotel_with_room(db_session, world.event.id)
    world.registration.is_participating = False
    await db_session.commit()

    resp = await client.post(
        f"/api/events/{world.event.id}/room-assignments/assign",
        headers=auth_headers(world.organizer_user),
        json={"employee_id": world.employee.id, "room_id": room.id},
    )
    assert resp.status_code == 400
    assert resp.json()["error"]["code"] == "invalid_employee_id"


async def test_room_capacity_cannot_drop_below_occupancy(
    client, world, auth_headers, db_session
):
    hotel, _room_type, room = await _make_hotel_with_room(db_session, world.event.id)
    second = await make_employee(db_session, team=world.team, site=world.site, code="NV-H02")
    db_session.add(
        Registration(
            event_id=world.event.id,
            employee_id=second.employee.id,
            status="submitted",
            is_participating=True,
            agreed_terms_at=utcnow(),
            terms_version="v1",
            submitted_at=utcnow(),
        )
    )
    await db_session.commit()
    for employee_id in (world.employee.id, second.employee.id):
        resp = await client.post(
            f"/api/events/{world.event.id}/room-assignments/assign",
            headers=auth_headers(world.organizer_user),
            json={"employee_id": employee_id, "room_id": room.id},
        )
        assert resp.status_code == 200

    resp = await client.patch(
        f"/api/events/{world.event.id}/hotels/{hotel.id}/rooms/{room.id}",
        headers=auth_headers(world.organizer_user),
        json={"capacity": 1},
    )
    assert resp.status_code == 409
    assert resp.json()["error"]["code"] == "capacity_below_assigned"


async def test_room_import_uses_hotel_code_when_room_numbers_repeat(
    client, world, auth_headers, db_session
):
    hotel_a = Hotel(event_id=world.event.id, code="HA", name="Hotel A")
    hotel_b = Hotel(event_id=world.event.id, code="HB", name="Hotel B")
    db_session.add_all([hotel_a, hotel_b])
    await db_session.flush()
    room_a = Room(hotel_id=hotel_a.id, room_number="101", capacity=2)
    room_b = Room(hotel_id=hotel_b.id, room_number="101", capacity=2)
    db_session.add_all([room_a, room_b])
    await db_session.commit()

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["employee_code", "hotel_code", "room_number"])
    sheet.append([world.employee.employee_code, "HB", "101"])
    payload = io.BytesIO()
    workbook.save(payload)

    resp = await client.post(
        f"/api/events/{world.event.id}/room-assignments/import",
        headers=auth_headers(world.organizer_user),
        files={"file": ("rooms.xlsx", payload.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200
    assert resp.json()["ok_rows"] == 1
    assignment = (
        await db_session.execute(
            select(RoomAssignment).where(RoomAssignment.employee_id == world.employee.id)
        )
    ).scalar_one()
    assert assignment.room_id == room_b.id


async def test_rooms_import_upserts_instead_of_failing_on_re_import(
    client, world, auth_headers, db_session
):
    """Re-importing a corrected room list must update in place — a blind insert
    hit uq_room_hotel_number and reported every unchanged row as a raw
    'UNIQUE constraint failed' error."""
    hotel, _room_type, _room = await _make_hotel_with_room(db_session, world.event.id)

    def _file(capacity: int, note: str) -> dict:
        wb = Workbook()
        ws = wb.active
        ws.append(["room_number", "capacity", "note"])
        ws.append(["101", capacity, note])
        ws.append(["102", 2, ""])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return {"file": ("rooms.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}

    first = await client.post(
        f"/api/events/{world.event.id}/hotels/{hotel.id}/rooms/import",
        headers=auth_headers(world.organizer_user), files=_file(2, "view bien"),
    )
    assert first.status_code == 200
    assert first.json() == {"ok_rows": 2, "error_rows": 0, "errors": []}

    second = await client.post(
        f"/api/events/{world.event.id}/hotels/{hotel.id}/rooms/import",
        headers=auth_headers(world.organizer_user), files=_file(3, "view nui"),
    )
    assert second.status_code == 200, second.text
    assert second.json()["error_rows"] == 0

    listed = await client.get(
        f"/api/events/{world.event.id}/hotels/{hotel.id}/rooms",
        headers=auth_headers(world.organizer_user),
    )
    rooms = listed.json()
    assert len(rooms) == 2, "re-import must not duplicate rooms"
    updated = next(r for r in rooms if r["room_number"] == "101")
    assert (updated["capacity"], updated["note"]) == (3, "view nui")


async def test_rooms_import_refuses_shrinking_below_occupancy(
    client, world, auth_headers, db_session
):
    hotel, _rt, room = await _make_hotel_with_room(db_session, world.event.id)
    roommate = await make_employee(
        db_session, team=world.team, site=world.site, code="NV030"
    )
    db_session.add_all([
        RoomAssignment(
            event_id=world.event.id, room_id=room.id, employee_id=world.employee.id,
            source="manual", assigned_at=utcnow(),
        ),
        RoomAssignment(
            event_id=world.event.id, room_id=room.id, employee_id=roommate.employee.id,
            source="manual", assigned_at=utcnow(),
        ),
    ])
    await db_session.commit()

    wb = Workbook()
    ws = wb.active
    ws.append(["room_number", "capacity"])
    ws.append(["101", 1])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    res = await client.post(
        f"/api/events/{world.event.id}/hotels/{hotel.id}/rooms/import",
        headers=auth_headers(world.organizer_user),
        files={"file": ("rooms.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert res.status_code == 200
    assert res.json()["error_rows"] == 1
    assert "đang có 2 người" in res.json()["errors"][0]["error"]


async def test_room_assignment_frozen_once_event_completed(
    client, world, auth_headers, db_session
):
    _hotel, _rt, room = await _make_hotel_with_room(db_session, world.event.id)
    world.event.status = EventStatus.event_completed
    await db_session.commit()

    res = await client.post(
        f"/api/events/{world.event.id}/room-assignments/assign",
        headers=auth_headers(world.organizer_user),
        json={"room_id": room.id, "employee_id": world.employee.id},
    )
    assert res.status_code == 400
    assert res.json()["error"]["code"] == "event_completed"
